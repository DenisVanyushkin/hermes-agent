"""The returned owner workbook may change decisions, but never role identity."""

from __future__ import annotations

import importlib.util
from pathlib import Path

import pytest
from openpyxl import load_workbook


MODULE_PATH = Path(__file__).resolve().parents[2] / "scripts" / "job_intel_shortlist_workbook.py"
spec = importlib.util.spec_from_file_location("job_intel_shortlist_workbook", MODULE_PATH)
assert spec and spec.loader
workbook = importlib.util.module_from_spec(spec)
spec.loader.exec_module(workbook)


def artifact() -> dict:
    base = {
        "company": "Example", "title": "Chief Product Officer", "location": "Remote",
        "source": "ATS", "url": "https://example.org/job/1",
        "rule_ids": ["seniority", "product"], "ruleset_version": "rf1-test",
        "description": "Build the product roadmap and lead the team.\nRequirements: P&L ownership.",
    }
    return {
        "release_id": "shortlist-2026-W39", "commit": "abc123", "run_ids": [1, 2],
        "ruleset_versions": ["rf1-test"], "census_sha256": "census",
        "items": [{**base, "vacancy_key": "yes-1", "role_fit_verdict": "accept"}],
        "rejected_sample": [{**base, "vacancy_key": "no-1", "role_fit_verdict": "reject"}],
    }


def test_full_vacancy_text_and_owner_roundtrip(tmp_path: Path) -> None:
    source = artifact()
    source["items"][0]["description"] = "A" * 40000 + "\nFinal requirement"
    path = tmp_path / "review.xlsx"
    projection = workbook.write_workbook(path, source, "attempt-001")
    book = load_workbook(path)
    assert book.sheetnames == ["release", "shortlist", "rejected_sample"]
    assert "".join(book["shortlist"].cell(2, column).value for column in (11, 12)) == source["items"][0]["description"]
    assert book["rejected_sample"]["K2"].value == source["rejected_sample"][0]["description"]
    book["shortlist"]["O2"] = "yes"
    book["shortlist"]["P2"] = "Good scope"
    book["rejected_sample"]["O2"] = "blocked_language"
    book.save(path)
    result = workbook.read_owner_workbook(path, source, "attempt-001")
    assert result["projection_sha256"] == projection
    assert result["decisions"] == {
        "yes-1": {"owner_decision": "yes", "owner_note": "Good scope", "sheet": "shortlist"},
        "no-1": {"owner_decision": "blocked_language", "owner_note": "", "sheet": "rejected_sample"},
    }


def test_model_summary_is_visible_and_frozen_in_owner_workbook(tmp_path: Path) -> None:
    source = artifact()
    source["items"][0]["summary"] = "Leads roadmap and owns P&L."
    source["items"][0]["summary_status"] = "ok"
    path = tmp_path / "review.xlsx"
    workbook.write_workbook(path, source, "attempt-001")
    book = load_workbook(path)
    assert book["shortlist"]["M1"].value == "summary"
    assert book["shortlist"]["M2"].value == "Leads roadmap and owns P&L."
    assert book["shortlist"]["N2"].value == "ok"
    book["shortlist"]["M2"] = "Invented scope"
    book.save(path)
    with pytest.raises(workbook.WorkbookContractError, match="frozen"):
        workbook.read_owner_workbook(path, source, "attempt-001")


@pytest.mark.parametrize("tamper", ["company", "description", "missing", "extra", "duplicate", "release", "formula", "extra_column"])
def test_frozen_projection_or_row_set_change_fails(tmp_path: Path, tamper: str) -> None:
    source = artifact()
    path = tmp_path / "review.xlsx"
    workbook.write_workbook(path, source, "attempt-001")
    book = load_workbook(path)
    if tamper == "company":
        book["shortlist"]["B2"] = "Other"
    elif tamper == "description":
        book["shortlist"]["K2"] = "Edited role"
    elif tamper == "missing":
        book["shortlist"].delete_rows(2)
    elif tamper == "extra":
        book["shortlist"].append(["new-role"])
    elif tamper == "duplicate":
        book["shortlist"].append(["yes-1"])
    elif tamper == "release":
        book["release"]["B2"] = "other-release"
    elif tamper == "formula":
        book["shortlist"]["P2"] = "=1+1"
    elif tamper == "extra_column":
        book["shortlist"]["Q2"] = "untracked edit"
    book.save(path)
    with pytest.raises(workbook.WorkbookContractError):
        workbook.read_owner_workbook(path, source, "attempt-001")


def test_source_text_starting_with_formula_marker_is_literal(tmp_path: Path) -> None:
    source = artifact()
    source["items"][0]["description"] = "=HYPERLINK(\"bad\")"
    path = tmp_path / "review.xlsx"
    workbook.write_workbook(path, source, "attempt-001")
    cell = load_workbook(path)["shortlist"]["K2"]
    assert cell.value == '=HYPERLINK("bad")'
    assert cell.data_type == "s"
