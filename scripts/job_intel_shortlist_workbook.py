"""Create and validate the self-contained weekly shortlist review workbook.

The frozen weekly artifact is the authority for role identities and text. The
owner may edit only decision and note cells; the returned workbook is checked
against the artifact before any decision can be imported.
"""

from __future__ import annotations

import hashlib
import json
from pathlib import Path
from typing import Any
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation


class WorkbookContractError(ValueError):
    """An owner workbook no longer matches its frozen weekly release."""


SCHEMA_VERSION = "1"
TEXT_CHUNK = 30_000
BASE_COLUMNS = (
    "vacancy_key", "company", "title", "location", "source", "url",
    "role_fit_verdict", "rule_ids", "ruleset_version", "run_id",
)
OWNER_COLUMNS = ("owner_decision", "owner_note")
SUMMARY_COLUMNS = ("summary", "summary_status")
RELEASE_FIELDS = (
    "release_id", "attempt_id", "schema_version", "projection_sha256",
    "commit", "run_ids", "ruleset_versions", "census_sha256",
)
DECISIONS = frozenset({"", "yes", "no", "blocked_language"})


def _text(value: Any) -> str:
    return "" if value is None else str(value)


def _row(item: dict[str, Any]) -> dict[str, str]:
    return {
        "vacancy_key": _text(item.get("vacancy_key")),
        "company": _text(item.get("company")),
        "title": _text(item.get("title")),
        "location": _text(item.get("location")),
        "source": _text(item.get("source")),
        "url": _text(item.get("url")),
        "role_fit_verdict": _text(item.get("role_fit_verdict")),
        "rule_ids": json.dumps(item.get("rule_ids") or [], ensure_ascii=False, separators=(",", ":")),
        "ruleset_version": _text(item.get("ruleset_version")),
        "run_id": _text(item.get("run_id")),
        "description": _text(item.get("description")),
        "summary": _text(item.get("summary")),
        "summary_status": _text(item.get("summary_status") or "unavailable"),
    }


def _sections(artifact: dict[str, Any]) -> dict[str, list[dict[str, str]]]:
    if artifact.get("artifact") not in (None, "job_intel_weekly_shortlist"):
        raise WorkbookContractError("not a weekly shortlist artifact")
    sections: dict[str, list[dict[str, str]]] = {}
    keys: set[str] = set()
    for sheet, field, verdicts in (
        ("shortlist", "items", {"accept"}),
        ("rejected_sample", "rejected_sample", {"reject", "blocked"}),
    ):
        items = artifact.get(field)
        if not isinstance(items, list):
            raise WorkbookContractError(f"artifact lacks {field}")
        rows = []
        for item in items:
            if not isinstance(item, dict):
                raise WorkbookContractError(f"invalid item in {field}")
            row = _row(item)
            key = row["vacancy_key"]
            if not key or key in keys:
                raise WorkbookContractError(f"empty or duplicate vacancy_key: {key!r}")
            if row["role_fit_verdict"] not in verdicts:
                raise WorkbookContractError(f"wrong verdict for {sheet}: {key}")
            if not row["description"].strip():
                raise WorkbookContractError(f"missing full description: {key}")
            keys.add(key)
            rows.append(row)
        sections[sheet] = rows
    return sections


def canonical_projection(artifact: dict[str, Any]) -> bytes:
    sections = _sections(artifact)
    projection = [
        [sheet, [row for row in sections[sheet]]]
        for sheet in ("shortlist", "rejected_sample")
    ]
    return json.dumps(projection, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")


def projection_sha256(artifact: dict[str, Any]) -> str:
    return hashlib.sha256(canonical_projection(artifact)).hexdigest()


def _description_columns(sections: dict[str, list[dict[str, str]]]) -> int:
    longest = max((len(row["description"]) for rows in sections.values() for row in rows), default=0)
    return max(2, (longest + TEXT_CHUNK - 1) // TEXT_CHUNK)


def _headers(description_columns: int) -> list[str]:
    return [*BASE_COLUMNS, *(f"description_{index}" for index in range(1, description_columns + 1)),
            *SUMMARY_COLUMNS, *OWNER_COLUMNS]


def _set_string(cell: Any, value: str) -> None:
    cell.value = value
    # Vacancy text is untrusted external data. A leading '=' must not become
    # a formula when the owner opens the workbook.
    cell.data_type = "s"


def _release_values(artifact: dict[str, Any], attempt_id: str, projection: str) -> dict[str, str]:
    if not artifact.get("release_id") or not attempt_id:
        raise WorkbookContractError("release_id and attempt_id are required")
    return {
        "release_id": _text(artifact["release_id"]),
        "attempt_id": attempt_id,
        "schema_version": SCHEMA_VERSION,
        "projection_sha256": projection,
        "commit": _text(artifact.get("commit")),
        "run_ids": json.dumps(artifact.get("run_ids") or [], separators=(",", ":")),
        "ruleset_versions": json.dumps(artifact.get("ruleset_versions") or [], separators=(",", ":")),
        "census_sha256": _text(artifact.get("census_sha256")),
    }


def write_workbook(path: Path, artifact: dict[str, Any], attempt_id: str) -> str:
    """Write one XLSX, returning the digest of all ordered frozen role fields."""
    sections = _sections(artifact)
    projection = projection_sha256(artifact)
    description_columns = _description_columns(sections)
    headers = _headers(description_columns)
    book = Workbook()
    release = book.active
    release.title = "release"
    for index, (field, value) in enumerate(_release_values(artifact, attempt_id, projection).items(), start=1):
        _set_string(release.cell(index, 1), field)
        _set_string(release.cell(index, 2), value)
        release.cell(index, 1).font = Font(bold=True)
    release.column_dimensions["A"].width = 24
    release.column_dimensions["B"].width = 76

    for sheet_name, rows in sections.items():
        sheet = book.create_sheet(sheet_name)
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="23395B")
        for row_number, row in enumerate(rows, start=2):
            values = [row[field] for field in BASE_COLUMNS]
            values += [row["description"][index * TEXT_CHUNK:(index + 1) * TEXT_CHUNK] for index in range(description_columns)]
            values += [row[field] for field in SUMMARY_COLUMNS]
            values += ["", ""]
            for column, value in enumerate(values, start=1):
                cell = sheet.cell(row_number, column)
                _set_string(cell, value)
                if column > len(BASE_COLUMNS) and column <= len(BASE_COLUMNS) + description_columns:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
            sheet.row_dimensions[row_number].height = 90
        sheet.freeze_panes = "B2"
        sheet.auto_filter.ref = f"A1:{sheet.cell(max(1, len(rows) + 1), len(headers)).coordinate}"
        sheet.column_dimensions["A"].width = 24
        sheet.column_dimensions["B"].width = 22
        sheet.column_dimensions["C"].width = 36
        for index in range(len(BASE_COLUMNS) + 1, len(BASE_COLUMNS) + description_columns + 1):
            sheet.column_dimensions[sheet.cell(1, index).column_letter].width = 70
        summary_column = len(BASE_COLUMNS) + description_columns + 1
        sheet.column_dimensions[sheet.cell(1, summary_column).column_letter].width = 70
        sheet.column_dimensions[sheet.cell(1, summary_column + 1).column_letter].width = 20
        owner_column = summary_column + len(SUMMARY_COLUMNS)
        sheet.column_dimensions[sheet.cell(1, owner_column).column_letter].width = 23
        sheet.column_dimensions[sheet.cell(1, owner_column + 1).column_letter].width = 54
        validation = DataValidation(type="list", formula1='"yes,no,blocked_language"', allow_blank=True)
        sheet.add_data_validation(validation)
        if rows:
            validation.add(f"{sheet.cell(2, owner_column).coordinate}:{sheet.cell(len(rows) + 1, owner_column).coordinate}")
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)
    return projection


def _cell_text(cell: Any) -> str:
    if cell.data_type == "f":
        raise WorkbookContractError(f"formula in returned workbook: {cell.coordinate}")
    return _text(cell.value)


def read_owner_workbook(path: Path, artifact: dict[str, Any], attempt_id: str) -> dict[str, Any]:
    """Return decisions only after every frozen workbook cell matches artifact."""
    sections = _sections(artifact)
    projection = projection_sha256(artifact)
    expected_release = _release_values(artifact, attempt_id, projection)
    description_columns = _description_columns(sections)
    headers = _headers(description_columns)
    try:
        book = load_workbook(path, read_only=False, data_only=False, keep_links=False)
    except (OSError, ValueError, KeyError, BadZipFile) as error:
        raise WorkbookContractError("cannot open owner workbook") from error
    if book.sheetnames != ["release", "shortlist", "rejected_sample"]:
        raise WorkbookContractError("workbook sheets changed")
    release = book["release"]
    if release.max_row != len(RELEASE_FIELDS):
        raise WorkbookContractError("release metadata row count changed")
    for index, field in enumerate(RELEASE_FIELDS, start=1):
        if _cell_text(release.cell(index, 1)) != field or _cell_text(release.cell(index, 2)) != expected_release[field]:
            raise WorkbookContractError(f"release metadata changed: {field}")
        if any(_cell_text(release.cell(index, column)) for column in range(3, release.max_column + 1)):
            raise WorkbookContractError(f"extra release metadata: {field}")
    decisions: dict[str, dict[str, str]] = {}
    for sheet_name, rows in sections.items():
        sheet = book[sheet_name]
        if sheet.max_row != len(rows) + 1:
            raise WorkbookContractError(f"row set changed: {sheet_name}")
        if [_cell_text(sheet.cell(1, column)) for column in range(1, len(headers) + 1)] != headers:
            raise WorkbookContractError(f"columns changed: {sheet_name}")
        if any(
            _cell_text(sheet.cell(row_number, column))
            for row_number in range(1, sheet.max_row + 1)
            for column in range(len(headers) + 1, sheet.max_column + 1)
        ):
            raise WorkbookContractError(f"extra columns: {sheet_name}")
        for row_number, expected in enumerate(rows, start=2):
            actual = {
                field: _cell_text(sheet.cell(row_number, column))
                for column, field in enumerate(BASE_COLUMNS, start=1)
            }
            actual["description"] = "".join(
                _cell_text(sheet.cell(row_number, len(BASE_COLUMNS) + index))
                for index in range(1, description_columns + 1)
            )
            for index, field in enumerate(SUMMARY_COLUMNS, start=1):
                actual[field] = _cell_text(sheet.cell(row_number, len(BASE_COLUMNS) + description_columns + index))
            if actual != expected:
                raise WorkbookContractError(f"frozen role changed: {sheet_name} row {row_number}")
            owner_column = len(BASE_COLUMNS) + description_columns + len(SUMMARY_COLUMNS) + 1
            decision = _cell_text(sheet.cell(row_number, owner_column)).strip().lower()
            note = _cell_text(sheet.cell(row_number, owner_column + 1))
            if decision not in DECISIONS:
                raise WorkbookContractError(f"invalid owner_decision: {expected['vacancy_key']}")
            if decision:
                decisions[expected["vacancy_key"]] = {
                    "owner_decision": decision, "owner_note": note, "sheet": sheet_name,
                }
    return {"release_id": expected_release["release_id"], "attempt_id": attempt_id,
            "projection_sha256": projection, "decisions": decisions}
